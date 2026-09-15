# -*- coding: utf-8 -*-
"""
世界数据 Word 文档构建器。

将世界的故事大纲、分集剧本、角色（含定妆图）、场景（含多视角图）、道具（含图片）
排版为一份 .docx 文档，供"导出 Word 文档"功能使用。

纯构建逻辑，不负责数据读取与文件上传——由 FileManager.export_world_docx 组装 data 后调用。
"""
import os
import re
from typing import Any, Dict, List, Optional, Tuple

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError as e:  # pragma: no cover
    raise ImportError("缺少依赖 python-docx，请先安装: pip install python-docx") from e

# 主题色
COLOR_TITLE = RGBColor(0xC0, 0x00, 0x00)   # 一级标题/角色名 红
COLOR_HEADING = RGBColor(0x1F, 0x3B, 0x63)  # 二级/三级标题、字幕 蓝
COLOR_GREY = RGBColor(0x80, 0x80, 0x80)     # 图注灰

# python-docx 支持的图片扩展名（其他格式 add_picture 会抛 ValueError）
_SUPPORTED_IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.gif', '.bmp', '.tif', '.tiff'}
# 单图嵌入上限：防止超大图撑爆内存（docx 需将图片解码进内存，zip 导出无此问题）
MAX_IMAGE_BYTES = 20 * 1024 * 1024


def _is_usable_image(image_path: Optional[str]) -> bool:
    """图片存在、扩展名受支持且不超过大小上限。"""
    if not image_path:
        return False
    try:
        if not os.path.exists(image_path):
            return False
        if os.path.splitext(image_path)[1].lower() not in _SUPPORTED_IMAGE_EXTS:
            return False
        return os.path.getsize(image_path) <= MAX_IMAGE_BYTES
    except OSError:
        return False


# ==================== 基础工具 ====================

def set_cn(style_or_run: Any, cn: str = "宋体") -> None:
    """设置中文字体（eastAsia）。注意必须走 get_or_add，直接 rFonts 可能为 None。"""
    el = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
    rPr = el.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), cn)


def _set_style_font(doc: Document, style_name: str, size: float, bold: bool = False,
                    name_cn: str = "宋体", color: Optional[RGBColor] = None) -> None:
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = color
    set_cn(style, name_cn)


def _add_image_or_placeholder(container: Any, image_path: Optional[str], width_cm: float,
                              url: str = "", caption: Optional[str] = None) -> None:
    """在 doc 或表格单元格中插入图片；缺失时降级为红色占位文字，不中断生成。"""
    if image_path and _is_usable_image(image_path):
        try:
            container.add_picture(image_path, width=Cm(width_cm))
            container.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception:
            # 损坏/无法解码的图片不中断导出，降级为占位文字
            p = container.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[图片无法嵌入: {url or '未知'}]")
            r.font.color.rgb = COLOR_TITLE
            set_cn(r)
    else:
        p = container.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[缺失图片: {url or '未知'}]")
        r.font.color.rgb = COLOR_TITLE
        set_cn(r)
    if caption:
        cp = container.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.font.size = Pt(9)
        r.font.color.rgb = COLOR_GREY
        set_cn(r)


def _shade_cell(cell: Any, hexcolor: str = "F2F2F2") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): hexcolor})
    tcPr.append(shd)


def _cell_image(cell: Any, image_path: Optional[str], width_cm: float, url: str = "") -> None:
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path and _is_usable_image(image_path):
        try:
            p.add_run().add_picture(image_path, width=Cm(width_cm))
            return
        except Exception:
            pass  # 损坏/无法解码的图片降级为占位文字
    r = p.add_run(f"[缺失图片: {url or '未知'}]")
    r.font.color.rgb = COLOR_TITLE
    set_cn(r)


def _desc_field(text: str, key: str) -> str:
    """从 'key：value\\nkey2：value2' 形式的描述文本中提取 key 对应的 value。"""
    m = re.search(re.escape(key) + r"[：:](.*?)(?=\n[A-Z一-鿿]*[：:]|$)", text or "", re.S)
    return m.group(1).strip() if m else ""


def _fmt_blocks(text: str) -> List[Tuple[Optional[str], str, str]]:
    """把 'key：val\\nkey：val' 拆成 [(key, sep, val)]，无 key 的段为 (None, '', seg)。"""
    out = []
    for seg in re.split(r"\n(?=[^：:\n]+[：:])", text or ""):
        m = re.match(r"([^：:]+)([：:])(.*)", seg.strip())
        if m:
            out.append((m.group(1).strip(), m.group(2), m.group(3).strip()))
        elif seg.strip():
            out.append((None, "", seg.strip()))
    return out


def _cn_run(p: Any, text: str, bold: bool = False, color: Optional[RGBColor] = None) -> Any:
    r = p.add_run(text)
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    set_cn(r)
    return r


# ==================== 封面 / 目录 ====================

def _extract_title(world: Optional[dict]) -> str:
    """优先取大纲首行《剧名》，没有则用世界名。"""
    outline = str((world or {}).get("story_outline") or "")
    m = re.search(r"《([^》]+)》", outline)
    if m:
        return m.group(1)
    return ((world or {}).get("name") or "世界剧本").strip()


def _render_cover(doc: Document, title: str) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("短剧剧本合集"); r.font.size = Pt(26); r.font.bold = True
    r.font.name = "Times New Roman"
    set_cn(r, "黑体")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = COLOR_TITLE
    r.font.name = "Times New Roman"
    set_cn(r, "黑体")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cn_run(p, "含：故事大纲 · 分集剧本 · 角色说明与定妆图 · 场景设计 · 道具设定").font.size = Pt(11)
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cn_run(p, "（本文档由世界数据自动导出生成）").font.size = Pt(10)
    doc.add_page_break()


def _render_toc(doc: Document) -> None:
    doc.add_heading("目录", level=1)
    for item in ["一、故事大纲", "二、角色说明（含角色图片）", "三、场景设计（含场景图片）",
                 "四、道具设定（含道具图片）", "五、完整剧本"]:
        p = doc.add_paragraph(item)
        for r in p.runs:
            set_cn(r)
    doc.add_page_break()


# ==================== 一、故事大纲 ====================

def _render_outline(doc: Document, world: Optional[dict]) -> None:
    doc.add_heading("一、故事大纲", level=1)

    info = [
        ("剧 本 名 称", _extract_title(world)),
        ("类  型", (world or {}).get("story_type") or "—"),
        ("视觉风格", (world or {}).get("visual_style") or "—"),
        ("时代背景", (world or {}).get("era_environment") or "—"),
        ("构图偏好", (world or {}).get("composition_preference") or "—"),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(info):
        c0, c1 = t.rows[i].cells
        c0.width, c1.width = Cm(3.5), Cm(12)
        r = c0.paragraphs[0].add_run(k); r.font.bold = True
        set_cn(r)
        _shade_cell(c0)
        r = c1.paragraphs[0].add_run(str(v))
        set_cn(r)
    doc.add_paragraph()

    outline = str((world or {}).get("story_outline") or "")
    if not outline.strip():
        _cn_run(doc.add_paragraph(), "（无故事大纲）")
        return

    doc.add_heading("故事大纲全文", level=2)
    for line in outline.splitlines():
        line_s = line.rstrip()
        if not line_s.strip():
            continue
        if line_s.startswith("# ") and not line_s.startswith("## "):
            h = doc.add_heading(line_s[2:], level=2)
        elif line_s.startswith("## "):
            h = doc.add_heading(line_s[3:], level=3)
        elif line_s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            txt = line_s[2:]
            parts = re.match(r"^([^\s：:]+)(：)(.*)$", txt)
            if parts:
                _cn_run(p, parts.group(1), bold=True)
                _cn_run(p, parts.group(2) + parts.group(3))
            else:
                _cn_run(p, txt)
        elif line_s.startswith("**") and line_s.endswith("**"):
            _cn_run(doc.add_paragraph(), line_s.strip("*"), bold=True)
        else:
            _cn_run(doc.add_paragraph(), line_s)
    doc.add_page_break()


# ==================== 二、角色说明 ====================

def _render_character(doc: Document, c: dict, images: Dict[str, str]) -> None:
    name = c.get("name", "未命名")
    doc.add_heading(f"角色：{name}", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lc, rc = table.rows[0].cells
    lc.width, rc.width = Cm(5.5), Cm(10)

    ref = c.get("reference_image") or ""
    _cell_image(lc, images.get(ref), 5.0, ref)
    cap = lc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(f"（{name} 定妆参考图）")
    r.font.size = Pt(9); r.font.color.rgb = COLOR_GREY
    set_cn(r)

    basic = [
        ("姓  名", name),
        ("年  龄", c.get("age", "—")),
        ("身  份", c.get("identity", "—")),
        ("默认配音", c.get("default_voice", "—")),
    ]
    bp = rc.paragraphs[0]
    for k, v in basic:
        _cn_run(bp, f"{k}：", bold=True)
        _cn_run(bp, str(v) + "    ")
    _cn_run(rc.add_paragraph(), "【外貌】", bold=True)
    _cn_run(rc.add_paragraph(), c.get("appearance", "—"))

    for field, label in (("personality", "性格"), ("behavior", "行为举止"), ("other_info", "背景与弧线")):
        _cn_run(rc.add_paragraph(), f"【{label}】", bold=True)
        for key, sep, val in _fmt_blocks(c.get(field, "")):
            pc = rc.add_paragraph()
            pc.paragraph_format.left_indent = Cm(0.5)
            if key is None:
                _cn_run(pc, val)
            else:
                _cn_run(pc, key + sep, bold=True)
                _cn_run(pc, val)
    doc.add_paragraph()


def _render_characters(doc: Document, characters: List[dict], images: Dict[str, str]) -> None:
    doc.add_heading("二、角色说明", level=1)
    p = doc.add_paragraph()
    _cn_run(p, f"本剧共设 {len(characters)} 名角色。每位角色的定妆参考图与完整人物小传如下。")
    for c in characters:
        _render_character(doc, c, images)
    doc.add_page_break()


# ==================== 三、场景设计 ====================

def _render_location(doc: Document, d: dict, images: Dict[str, str]) -> None:
    name = d.get("name", "未命名").replace("_", " / ")
    doc.add_heading(f"场景：{name}", level=2)
    ref = d.get("reference_image") or ""
    _add_image_or_placeholder(doc, images.get(ref), 14.0, ref, caption=f"{name}（主视角）")

    raw_angles = d.get("reference_images")
    angles = sorted((a for a in raw_angles if isinstance(a, dict)),
                    key=lambda x: (x.get("label") or "")) if isinstance(raw_angles, list) else []
    if angles:
        t = doc.add_table(rows=2, cols=len(angles))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, a in enumerate(angles):
            url = a.get("url") or ""
            _cell_image(t.rows[0].cells[j], images.get(url), 4.8, url)
            cp = t.rows[1].cells[j].paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cp.add_run(a.get("label") or "")
            r.font.size = Pt(9); r.font.color.rgb = COLOR_GREY
            set_cn(r)
    doc.add_paragraph()

    desc = str(d.get("description") or "")
    for key in ["类型", "功能", "规模", "布局", "特征", "环境", "氛围", "剧情作用"]:
        val = _desc_field(desc, key)
        if val:
            p = doc.add_paragraph()
            _cn_run(p, f"【{key}】", bold=True)
            _cn_run(p, val)


def _render_locations(doc: Document, locations: List[dict], images: Dict[str, str]) -> None:
    doc.add_heading("三、场景设计", level=1)
    p = doc.add_paragraph()
    _cn_run(p, f"本剧共设 {len(locations)} 个主要场景，每个场景均提供主视角图及环绕参考图。")
    for d in locations:
        _render_location(doc, d, images)
    doc.add_page_break()


# ==================== 四、道具设定 ====================

def _render_prop(doc: Document, d: dict, images: Dict[str, str]) -> None:
    name = d.get("name", "未命名")
    doc.add_heading(f"道具：{name}", level=2)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    lc, rc = t.rows[0].cells
    lc.width, rc.width = Cm(5.5), Cm(10)
    ref = d.get("reference_image") or ""
    _cell_image(lc, images.get(ref), 4.5, ref)
    _cn_run(rc.paragraphs[0], f"类型：{d.get('type', '—')}", bold=True)
    desc = str(d.get("description") or "")
    for key in ["外观", "规格", "功能", "特殊性", "背景", "剧情作用", "象征意义"]:
        val = _desc_field(desc, key)
        if val:
            p = rc.add_paragraph()
            _cn_run(p, f"【{key}】", bold=True)
            _cn_run(p, val)
    doc.add_paragraph()


def _render_props(doc: Document, props: List[dict], images: Dict[str, str]) -> None:
    doc.add_heading("四、道具设定", level=1)
    p = doc.add_paragraph()
    _cn_run(p, f"本剧共设 {len(props)} 件关键道具。")
    for d in props:
        _render_prop(doc, d, images)
    doc.add_page_break()


# ==================== 五、完整剧本 ====================

def _render_script(doc: Document, s: dict) -> None:
    """渲染单集剧本 markdown 内容：场标题/角色台词/字幕/转场/片尾。"""
    content = str(s.get("content") or "")
    # 集标题：优先取内容首行 H1，其次 title + 集数
    ep_line = next((l for l in content.splitlines() if l.startswith("# ")), None)
    if ep_line:
        title = ep_line[2:].strip()
    else:
        title = s.get("title") or "剧本"
        if s.get("episode_number"):
            title = f"{title}（第{s['episode_number']}集）"
    doc.add_heading(title, level=2)

    for line in content.splitlines():
        s_ = line.strip()
        if not s_ or s_ == "---":
            continue
        if s_.startswith("# "):
            continue  # 集标题已处理
        if s_.startswith("## "):
            h = doc.add_heading(s_[3:], level=3)
            for r in h.runs:
                set_cn(r, "黑体")
            continue
        if s_.startswith("【") and s_.endswith("】"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _cn_run(p, s_, bold=True, color=COLOR_HEADING)
            continue
        if s_.startswith("**（") and s_.endswith("）**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _cn_run(p, s_.strip("*"), bold=True)
            continue
        if s_.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            m = re.match(r"^-(.+?)[：:](.*)$", s_[2:])
            if m:
                _cn_run(p, m.group(1), bold=True)
                _cn_run(p, "：" + m.group(2))
            else:
                _cn_run(p, s_[2:])
            continue
        m = re.match(r"^\*\*(.+?)\*\*", s_)
        if m and ("：" in s_ or ":" in s_):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            _cn_run(p, m.group(1), bold=True, color=COLOR_TITLE)
            _cn_run(p, s_[len(m.group(0)):])
            continue
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.74)
        _cn_run(p, s_)


def _render_scripts(doc: Document, scripts: List[dict]) -> None:
    doc.add_heading("五、完整剧本", level=1)
    ordered = sorted(scripts, key=lambda x: (x.get("episode_number") is None, x.get("episode_number")))
    for i, s in enumerate(ordered):
        _render_script(doc, s)
        if i < len(ordered) - 1:
            doc.add_page_break()


# ==================== 入口 ====================

def build_world_docx(data: Dict[str, Any], out_path: str) -> None:
    """
    根据世界数据生成 Word 文档。

    Args:
        data: {
            "world": dict | None,            # worlds/world_<id>.json 原始数据
            "characters": [dict, ...],       # 角色 JSON 列表
            "locations": [dict, ...],        # 场景 JSON 列表
            "props": [dict, ...],            # 道具 JSON 列表
            "scripts": [dict, ...],          # 剧本 JSON 列表（含 title/episode_number/content）
            "images": {url: 本地图片路径},   # reference_image / reference_images[].url → 本地路径
        }
        out_path: docx 输出路径
    """
    world = data.get("world")
    characters = data.get("characters") or []
    locations = data.get("locations") or []
    props = data.get("props") or []
    scripts = data.get("scripts") or []
    images = data.get("images") or {}

    doc = Document()

    # 全局字体：正文宋体 12pt，标题黑体
    _set_style_font(doc, "Normal", size=12)
    _set_style_font(doc, "Heading 1", size=22, bold=True, name_cn="黑体", color=COLOR_TITLE)
    _set_style_font(doc, "Heading 2", size=16, bold=True, name_cn="黑体", color=COLOR_HEADING)
    _set_style_font(doc, "Heading 3", size=14, bold=True, name_cn="黑体", color=COLOR_HEADING)

    for section in doc.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    _render_cover(doc, _extract_title(world))
    _render_toc(doc)
    _render_outline(doc, world)
    _render_characters(doc, characters, images)
    _render_locations(doc, locations, images)
    _render_props(doc, props, images)
    if scripts:
        _render_scripts(doc, scripts)
    elif not (characters or locations or props):
        _cn_run(doc.add_paragraph(), "（本世界暂无剧本、角色、场景或道具数据）")

    doc.save(out_path)
